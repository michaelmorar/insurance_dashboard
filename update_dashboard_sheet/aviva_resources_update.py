from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import re
import shutil
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl


warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    category=UserWarning,
    module="openpyxl.worksheet._reader",
)


WORKDIR = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = WORKDIR / "inputs" / "Aviva Resource Tracker_May 2025.xlsm"
DEFAULT_DESTINATION = WORKDIR / "outputs" / "Aviva Delivery Dashboard - 2026.xlsx"

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
NS = {"a": MAIN_NS, "r": REL_NS}
ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)
ET.register_namespace("x14ac", "http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac")
ET.register_namespace("mc", "http://schemas.openxmlformats.org/markup-compatibility/2006")
ET.register_namespace("xr", "http://schemas.microsoft.com/office/spreadsheetml/2014/revision")
ET.register_namespace("xr2", "http://schemas.microsoft.com/office/spreadsheetml/2015/revision2")
ET.register_namespace("xr3", "http://schemas.microsoft.com/office/spreadsheetml/2016/revision3")


def previous_month(today: dt.date | None = None) -> dt.date:
    today = today or dt.date.today()
    first_this_month = today.replace(day=1)
    previous_last_day = first_this_month - dt.timedelta(days=1)
    return previous_last_day.replace(day=1)


def parse_month(value: str) -> tuple[dt.datetime, dt.datetime, str]:
    if value.lower() == "previous":
        start_date = previous_month()
    else:
        try:
            start_date = dt.datetime.strptime(value, "%Y-%m").date().replace(day=1)
        except ValueError as exc:
            raise argparse.ArgumentTypeError("month must be YYYY-MM or 'previous'") from exc
    if start_date.month == 12:
        end_date = start_date.replace(year=start_date.year + 1, month=1)
    else:
        end_date = start_date.replace(month=start_date.month + 1)
    label = start_date.strftime("%b%y")
    return (
        dt.datetime.combine(start_date, dt.time()),
        dt.datetime.combine(end_date, dt.time()),
        label,
    )


def default_output_path(destination: Path, month_label: str) -> Path:
    return destination.with_name(f"{destination.stem}_updated_{month_label}{destination.suffix}")


def default_audit_path(month_label: str) -> Path:
    return WORKDIR / f"{month_label.lower()}_new_starters_audit.csv"


def default_duplicate_audit_path(month_label: str) -> Path:
    return WORKDIR / f"{month_label.lower()}_destination_duplicate_names.csv"


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def excel_serial(value: dt.datetime | dt.date) -> float:
    if isinstance(value, dt.date) and not isinstance(value, dt.datetime):
        value = dt.datetime.combine(value, dt.time())
    return (value - dt.datetime(1899, 12, 30)).days + (
        value.hour * 3600 + value.minute * 60 + value.second
    ) / 86400


def country_for_city(city: str, onshore_offshore: str) -> str:
    city_l = city.lower().strip()
    if city_l in {"bangalore", "bengaluru", "gurgaon", "gurugram", "pune", "hyderabad", "chennai"}:
        return "India"
    if city_l in {"london", "edinburgh", "manchester", "york", "sheffield", "glasgow", "bristol", "norwich", "perth"}:
        return "UK"
    if "offshore" in onshore_offshore.lower():
        return "India"
    return ""


def business_unit(area: str, project: str) -> str:
    area_l = area.lower().strip()
    if area_l in {"ukgi cio", "dlg", "general insurance"}:
        return "General Insurance"
    if area_l in {"life & pension", "iwr cio", "wealth"}:
        return "Insurance, Wealth & Retirement"
    if area_l == "aviva international":
        return "Cross-Group functions"
    if "river" in project.lower():
        return "General Insurance"
    return ""


def project_name(detail: str, sub_area: str) -> str:
    d = norm(detail)
    s = norm(sub_area)
    text = f"{d} {s}".lower()
    if "supplier data" in text:
        return "Claims Supplier Data"
    if "consumer duty" in text:
        return "New Consumer Duty"
    if "continuous improvement and propositions" in text:
        return "Continuous Improvement and Propositions program"
    if "motor acceleration" in text:
        return "PPUD - Motor Acceleration (River)"
    if "cuo" in text and "process integration" in text:
        return "PPUD - CUO & PL Process Integration"
    if "river finance integration" in text:
        return "Project River"
    if s and s.lower() not in {"ukgi", "dlg", "gi", "claims", "project manager"}:
        return s
    cleaned = re.sub(r"^(priority\s+)?\d+(/\d+)?\s*x\s+", "", d, flags=re.I)
    cleaned = re.sub(r"^(project manager|business analysis|business analyst|pm)\s*-\s*", "", cleaned, flags=re.I)
    cleaned = re.sub(r"\s+-\s+Aviva CS.*$", "", cleaned)
    return cleaned or s


def project_type(project: str) -> str:
    p = project.lower()
    if "consumer duty" in p:
        return "Regulatory change"
    if "supplier data" in p:
        return "Strategic project"
    if "continuous improvement" in p:
        return "Digital transformation"
    if "river" in p or "ppud" in p:
        return "Data migration"
    return ""


def role_name(role: str, detail: str) -> str:
    r = norm(role)
    if r.lower() == "agile delivery specialist" and "product owner" in detail.lower():
        return "Product Owner / Manager"
    return r


def summary_text(detail: str, skills: str, comments: str, source_row: int) -> str:
    parts = []
    if detail:
        parts.append(norm(detail))
    if skills:
        parts.append(f"Specialist skills: {norm(skills)}")
    if comments:
        parts.append(f"Source comments: {norm(comments)}")
    parts.append(f"Source: Change Sourcing row {source_row}")
    return "\n".join(parts)


def stakeholder_text(hiring_manager: str, resource_manager: str) -> str:
    parts = []
    if hiring_manager:
        parts.append(f"Hiring Manager: {norm(hiring_manager)}")
    if resource_manager:
        parts.append(f"Aviva Resource Manager: {norm(resource_manager)}")
    return "; ".join(parts)


def row_key_tokens(project: str, summary: str) -> set[str]:
    stop = {
        "and",
        "for",
        "the",
        "with",
        "project",
        "program",
        "programme",
        "manager",
        "business",
        "analysis",
        "analyst",
        "aviva",
        "capco",
        "source",
        "change",
        "digital",
        "product",
        "owner",
        "management",
    }
    words = re.findall(r"[a-z0-9]+", f"{project} {summary}".lower())
    return {w for w in words if len(w) > 3 and w not in stop}


def project_key_tokens(project: str) -> set[str]:
    return row_key_tokens(project, "")


def read_existing_rows(destination: Path, destination_sheet: str) -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(destination, read_only=False, data_only=True)
    ws = wb[destination_sheet]
    rows = []
    for row_number in range(2, ws.max_row + 1):
        name = norm(ws.cell(row_number, 1).value)
        if not name:
            continue
        rows.append(
            {
                "row": row_number,
                "name": name,
                "project": norm(ws.cell(row_number, 3).value),
                "summary": norm(ws.cell(row_number, 7).value),
            }
        )
    return rows


def existing_match(
    name: str,
    project: str,
    summary: str,
    existing_rows: list[dict[str, object]],
    destination_sheet: str,
) -> tuple[bool, str]:
    name_l = name.lower()
    for existing in existing_rows:
        existing_name = str(existing["name"]).lower()
        same_name = existing_name == name_l
        caroline_typo = name_l == "caroline pham" and existing_name == "carolone pham"
        if not (same_name or caroline_typo):
            continue
        return True, f"name already present on {destination_sheet} row {existing['row']}"
    return False, ""


def extract_source_rows(
    source: Path,
    source_sheet: str,
    existing_rows: list[dict[str, object]],
    start: dt.datetime,
    end: dt.datetime,
    destination_sheet: str,
    status_value: str,
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_vba=True)
    ws = wb[source_sheet]
    to_add = []
    audit = []
    seen_names = {str(row["name"]).lower() for row in existing_rows}

    for row_number, values in enumerate(ws.iter_rows(min_row=2, max_col=39, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            if row_number > 1500:
                break
            continue
        status = values[14]
        actual_start = values[36]
        if status != status_value:
            continue
        if not isinstance(actual_start, dt.datetime) or not (start <= actual_start < end):
            continue

        name = norm(values[15])
        detail = norm(values[3])
        aviva_area = norm(values[4])
        sub_area = norm(values[5])
        project = project_name(detail, sub_area)
        summary = summary_text(detail, norm(values[9]), norm(values[38]), row_number)
        add = True
        reason = "added"

        if not name or name.lower() == "tbc":
            add = False
            reason = "skipped: candidate name is TBC"
        elif name.lower() in seen_names:
            add = False
            reason = f"skipped: name already present in {destination_sheet}"
        else:
            matched, matched_reason = existing_match(name, project, summary, existing_rows, destination_sheet)
            if matched:
                add = False
                reason = f"skipped: {matched_reason}"

        record = {
            "source_row": row_number,
            "name": name,
            "actual_start": actual_start,
            "project": project,
            "detail": detail,
            "business_unit": business_unit(aviva_area, project),
            "project_type": project_type(project),
            "change_sourcing": "Yes" if norm(values[0]).upper() == "YES" else "No",
            "active": "Yes",
            "summary": summary,
            "role": role_name(norm(values[6]), detail),
            "stakeholders": stakeholder_text(norm(values[11]), norm(values[12])),
            "capco_line_manager": norm(values[13]),
            "city": norm(values[33]),
            "country": country_for_city(norm(values[33]), norm(values[10])),
            "aviva_role_location": norm(values[33]) or ("India" if "offshore" in norm(values[10]).lower() else ""),
            "action": reason,
        }
        audit.append(record)
        if add:
            to_add.append(record)
            seen_names.add(name.lower())

    return to_add, audit


def shared_strings(zip_file: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zip_file.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    values = []
    for si in root.findall("a:si", NS):
        texts = [node.text or "" for node in si.findall(".//a:t", NS)]
        values.append("".join(texts))
    return values


def sheet_path_for(zip_file: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = ET.fromstring(zip_file.read("xl/workbook.xml"))
    rels = ET.fromstring(zip_file.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
    for sheet in workbook.find("a:sheets", NS):
        if sheet.attrib["name"] == sheet_name:
            rid = sheet.attrib[f"{{{REL_NS}}}id"]
            target = rid_to_target[rid].lstrip("/")
            return target if target.startswith("xl/") else f"xl/{target}"
    raise ValueError(f"Sheet not found: {sheet_name}")


def cell_ref(column: int, row: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def set_cell(cell: ET.Element, value: object, style: str | None = None) -> None:
    cell.attrib.pop("t", None)
    for child in list(cell):
        cell.remove(child)
    if style is not None:
        cell.attrib["s"] = style
    if value in (None, ""):
        return
    if isinstance(value, (dt.datetime, dt.date)):
        v = ET.SubElement(cell, f"{{{MAIN_NS}}}v")
        v.text = str(int(excel_serial(value)))
        return
    if isinstance(value, (int, float)):
        v = ET.SubElement(cell, f"{{{MAIN_NS}}}v")
        v.text = str(value)
        return
    cell.attrib["t"] = "inlineStr"
    is_node = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
    t_node = ET.SubElement(is_node, f"{{{MAIN_NS}}}t")
    text = str(value)
    if text != text.strip() or "\n" in text:
        t_node.attrib[f"{{{XML_NS}}}space"] = "preserve"
    t_node.text = text


def read_inline_or_shared(cell: ET.Element, strings: list[str]) -> str:
    if cell.attrib.get("t") == "s":
        node = cell.find("a:v", NS)
        if node is None or node.text is None:
            return ""
        index = int(node.text)
        return strings[index] if 0 <= index < len(strings) else ""
    if cell.attrib.get("t") == "inlineStr":
        texts = [node.text or "" for node in cell.findall(".//a:t", NS)]
        return "".join(texts)
    node = cell.find("a:v", NS)
    return node.text if node is not None and node.text is not None else ""


def clear_row_contents(row: ET.Element, max_column: int = 19) -> None:
    row_number = int(row.attrib["r"])
    existing = {}
    for cell in list(row.findall("a:c", NS)):
        col = re.sub(r"\d+", "", cell.attrib["r"])
        existing[col] = cell.attrib.get("s")
        row.remove(cell)
    for index in range(1, max_column + 1):
        ref = cell_ref(index, row_number)
        col = re.sub(r"\d+", "", ref)
        attrs = {"r": ref}
        if existing.get(col) is not None:
            attrs["s"] = existing[col]
        ET.SubElement(row, f"{{{MAIN_NS}}}c", attrs)


def row_has_content(row: ET.Element, strings: list[str], max_column: int = 19) -> bool:
    row_number = int(row.attrib["r"])
    for index in range(1, max_column + 1):
        ref = cell_ref(index, row_number)
        for cell in row.findall("a:c", NS):
            if cell.attrib.get("r") == ref and norm(read_inline_or_shared(cell, strings)):
                return True
    return False


def correct_known_typos(rows: dict[int, ET.Element], strings: list[str]) -> None:
    for row_number, row in rows.items():
        for cell in row.findall("a:c", NS):
            if cell.attrib.get("r") != f"A{row_number}":
                continue
            if norm(read_inline_or_shared(cell, strings)).lower() == "carolone pham":
                set_cell(cell, "Caroline Pham", cell.attrib.get("s"))


def dedupe_destination_rows(rows: dict[int, ET.Element], strings: list[str]) -> list[dict[str, object]]:
    seen: dict[str, int] = {}
    duplicates = []
    for row_number in sorted(rows):
        if row_number == 1:
            continue
        row = rows[row_number]
        name_cell = None
        for cell in row.findall("a:c", NS):
            if cell.attrib.get("r") == f"A{row_number}":
                name_cell = cell
                break
        if name_cell is None:
            continue
        name = norm(read_inline_or_shared(name_cell, strings))
        if not name:
            continue
        key = name.lower()
        if key in seen:
            duplicates.append({"name": name, "kept_row": seen[key], "cleared_row": row_number})
            clear_row_contents(row)
        else:
            seen[key] = row_number
    return duplicates


def patch_dashboard(
    destination: Path,
    output: Path,
    destination_sheet: str,
    to_add: list[dict[str, object]],
    dedupe_names: bool,
) -> list[dict[str, object]]:
    temp_output = output.with_suffix(".tmp.xlsx")
    shutil.copyfile(destination, temp_output)
    with zipfile.ZipFile(temp_output, "r") as zip_file:
        path = sheet_path_for(zip_file, destination_sheet)
        original_xml = zip_file.read(path)
        for prefix, uri in re.findall(r'xmlns(?::(\w+))?=[\'\"]([^\'\"]+)[\'\"]', original_xml.decode("utf-8", errors="ignore")):
            ET.register_namespace(prefix, uri)

        strings = shared_strings(zip_file)
        root = ET.fromstring(original_xml)
        sheet_data = root.find("a:sheetData", NS)
        rows = {int(row.attrib["r"]): row for row in sheet_data.findall("a:row", NS)}
        correct_known_typos(rows, strings)
        duplicate_rows = dedupe_destination_rows(rows, strings) if dedupe_names else []
        last_used = max(row_number for row_number, row in rows.items() if row_has_content(row, strings))
        template_row = rows.get(last_used + 1)
        if template_row is None:
            template_row = rows[last_used]
        template_styles = {}
        for cell in template_row.findall("a:c", NS):
            ref = cell.attrib["r"]
            col = re.sub(r"\d+", "", ref)
            template_styles[col] = cell.attrib.get("s")

        columns = [
            "name",
            "business_unit",
            "project",
            "project_type",
            "change_sourcing",
            "active",
            "summary",
            "role",
            "",
            "stakeholders",
            "capco_line_manager",
            "",
            "city",
            "country",
            "aviva_role_location",
            "actual_start",
            "",
            "",
            "",
        ]

        for offset, record in enumerate(to_add, start=1):
            row_number = last_used + offset
            if row_number in rows:
                row = rows[row_number]
                for child in list(row):
                    row.remove(child)
            else:
                row = copy.deepcopy(template_row)
                row.attrib["r"] = str(row_number)
                sheet_data.append(row)
                rows[row_number] = row
                for child in list(row):
                    row.remove(child)
            row.attrib["spans"] = "1:19"

            for index, key in enumerate(columns, start=1):
                col = re.sub(r"\d+", "", cell_ref(index, row_number))
                cell = ET.SubElement(row, f"{{{MAIN_NS}}}c", {"r": cell_ref(index, row_number)})
                style = template_styles.get(col)
                if style is not None:
                    cell.attrib["s"] = style
                value = record.get(key, "") if key else ""
                if key == "actual_start":
                    style = template_styles.get("P", style)
                set_cell(cell, value, style)

        patched_sheet = ET.tostring(root, encoding="utf-8", xml_declaration=True)

        orig_start = original_xml.find(b"<worksheet")
        orig_end = original_xml.find(b">", orig_start)
        original_worksheet_tag = original_xml[orig_start : orig_end + 1]

        namespaces = re.findall(r'xmlns(?::(\w+))?=[\'\"]([^\'\"]+)[\'\"]', original_xml.decode("utf-8", errors="ignore"))
        for prefix, uri in namespaces:
            search_str = f"xmlns:{prefix}=" if prefix else "xmlns="
            if search_str not in original_worksheet_tag.decode("utf-8", errors="ignore"):
                decl = f' xmlns:{prefix}="{uri}"' if prefix else f' xmlns="{uri}"'
                original_worksheet_tag = original_worksheet_tag[:-1] + decl.encode("utf-8") + b">"

        patch_start = patched_sheet.find(b"<worksheet")
        patch_end = patched_sheet.find(b">", patch_start)

        if orig_start != -1 and patch_start != -1:
            patched_sheet = (
                patched_sheet[:patch_start]
                + original_worksheet_tag
                + patched_sheet[patch_end + 1 :]
            )

    with zipfile.ZipFile(temp_output, "r") as source_zip, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target_zip:
        for item in source_zip.infolist():
            if item.filename == path:
                continue
            target_zip.writestr(item, source_zip.read(item.filename))
        target_zip.writestr(path, patched_sheet)
    temp_output.unlink()
    return duplicate_rows


def write_audit(path: Path, audit_rows: list[dict[str, object]]) -> None:
    fieldnames = [
        "source_row",
        "name",
        "actual_start",
        "project",
        "business_unit",
        "project_type",
        "role",
        "city",
        "country",
        "action",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in audit_rows:
            item = row.copy()
            if isinstance(item["actual_start"], dt.datetime):
                item["actual_start"] = item["actual_start"].strftime("%Y-%m-%d")
            writer.writerow({field: item.get(field, "") for field in fieldnames})


def write_duplicate_audit(path: Path, duplicate_rows: list[dict[str, object]]) -> None:
    fieldnames = ["name", "kept_row", "cleared_row"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(duplicate_rows)


def duplicate_names_from_existing(existing_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    seen: dict[str, int] = {}
    duplicates = []
    for row in existing_rows:
        name = str(row["name"])
        key = name.lower()
        if key in seen:
            duplicates.append({"name": name, "kept_row": seen[key], "cleared_row": row["row"]})
        else:
            seen[key] = int(row["row"])
    return duplicates


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Add previous-month Change Sourcing starters to an Aviva dashboard monthly DATA tab."
    )
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="Path to the Resource Tracker workbook.")
    parser.add_argument("--destination", type=Path, default=DEFAULT_DESTINATION, help="Path to the dashboard workbook.")
    parser.add_argument("--month", default="previous", help="Month to process, as YYYY-MM, or 'previous'.")
    parser.add_argument("--source-sheet", default="Change Sourcing", help="Source worksheet name.")
    parser.add_argument("--destination-sheet", help="Destination worksheet name. Defaults to DATA_<Month>, e.g. DATA_Apr26.")
    parser.add_argument("--status", default="12. Started at Aviva", help="Source status value to include.")
    parser.add_argument("--output", type=Path, help="Path for the updated workbook.")
    parser.add_argument("--audit", type=Path, help="Path for the starter audit CSV.")
    parser.add_argument("--duplicate-audit", type=Path, help="Path for the duplicate-name audit CSV.")
    parser.add_argument("--dry-run", action="store_true", help="Write audits only; do not create an updated workbook.")
    parser.add_argument(
        "--keep-duplicate-names",
        action="store_true",
        help="Do not clear duplicate names already present in the destination sheet.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    start, end, month_label = parse_month(args.month)
    destination_sheet = args.destination_sheet or f"DATA_{month_label}"
    output = args.output or default_output_path(args.destination, month_label)
    audit = args.audit or default_audit_path(month_label)
    duplicate_audit = args.duplicate_audit or default_duplicate_audit_path(month_label)

    existing_rows = read_existing_rows(args.destination, destination_sheet)
    to_add, audit_rows = extract_source_rows(
        args.source,
        args.source_sheet,
        existing_rows,
        start,
        end,
        destination_sheet,
        args.status,
    )
    if args.dry_run:
        duplicate_rows = duplicate_names_from_existing(existing_rows)
    else:
        duplicate_rows = patch_dashboard(
            args.destination,
            output,
            destination_sheet,
            to_add,
            dedupe_names=not args.keep_duplicate_names,
        )
    write_audit(audit, audit_rows)
    write_duplicate_audit(duplicate_audit, duplicate_rows)

    print(f"Processed month: {month_label} ({start:%Y-%m-%d} to {(end - dt.timedelta(days=1)):%Y-%m-%d})")
    print(f"Destination sheet: {destination_sheet}")
    print(f"Extracted starters: {len(audit_rows)}")
    print(f"Rows to add: {len(to_add)}")
    print(f"Duplicate destination names: {len(duplicate_rows)}")
    if args.dry_run:
        print("Dry run: no workbook written")
    else:
        print(f"Output workbook: {output}")
    print(f"Audit CSV: {audit}")
    print(f"Duplicate audit CSV: {duplicate_audit}")
    for item in audit_rows:
        print(f"{item['action']}: {item['name']} ({item['actual_start']:%Y-%m-%d})")


if __name__ == "__main__":
    main()
