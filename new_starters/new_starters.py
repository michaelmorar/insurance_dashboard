"""
compose_aviva_new_starter_emails.py

Purpose:
    Read the Aviva resource tracker, find new starters from the
    "Change Sourcing" tab, and compose deterministic welcome emails.

Input:
    aviva_resource_tracker.xlsm

Definition of a new starter:
    - Status contains "Started at Aviva"
    - Actual Start Date is within the selected month

Output:
    - One .eml email file per new starter
    - A CSV audit file listing the starters found

No OpenAI API key required.
No email is sent automatically.

Install:
    pip install openpyxl python-dateutil

Example:
    python compose_aviva_new_starter_emails.py \
      --workbook "aviva_resource_tracker.xlsm" \
      --month 2026-07 \
      --review-recipient "michael.morar@aviva.com" \
      --output-dir "new_starter_emails"
"""

from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any

from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook


SURVEY_LINK = (
    "https://forms.office.com/Pages/DesignPageV2.aspx?"
    "subpage=design&id=3Dmsm6unIEqai4hcYjW2xWcc3zskeKNGlTF9sOV2MWNUNlZSWTA4OEJTMlFHR0VLVTJGNVRFOTBNMi4u&analysis=false"
)

SHEET_NAME = "Change Sourcing"


@dataclass(frozen=True)
class NewStarter:
    candidate_name: str
    role_title: str | None
    aviva_business_area: str | None
    aviva_sub_business_area: str | None
    capco_lead: str | None
    hiring_manager: str | None
    status: str | None
    actual_start_date: date
    target_start_date: date | None
    recipient_email: str | None = None


def clean_text(value: Any) -> str | None:
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    return " ".join(text.split())


def normalise_header(value: Any) -> str:
    text = clean_text(value) or ""
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.casefold().strip()


def normalise_status(value: Any) -> str:
    return (clean_text(value) or "").casefold()


def parse_month(value: str | None) -> tuple[date, date]:
    """
    Parse YYYY-MM into an inclusive start date and exclusive end date.
    Defaults to the current month.
    """
    if value:
        start = datetime.strptime(value, "%Y-%m").date().replace(day=1)
    else:
        today = date.today()
        start = today.replace(day=1)

    end = start + relativedelta(months=1)
    return start, end


def parse_excel_date(value: Any) -> date | None:
    """
    Convert common Excel/openpyxl date values into a date.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, str):
        text = value.strip()

        if not text:
            return None

        accepted_formats = [
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d.%m.%y",
            "%d.%m.%Y",
            "%d-%m-%Y",
            "%d %b %Y",
            "%d %B %Y",
        ]

        for fmt in accepted_formats:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass

    return None


def find_header_map(ws) -> dict[str, int]:
    """
    Build a map from normalised column header to 1-based column number.
    Assumes headers are on row 1.
    """
    header_map: dict[str, int] = {}

    for cell in ws[1]:
        header = normalise_header(cell.value)

        if header:
            header_map[header] = cell.column

    return header_map


def require_column(header_map: dict[str, int], possible_names: list[str]) -> int:
    for name in possible_names:
        key = normalise_header(name)

        if key in header_map:
            return header_map[key]

    raise KeyError(
        f"Could not find required column. Tried: {possible_names}. "
        f"Available columns: {sorted(header_map.keys())}"
    )


def optional_column(header_map: dict[str, int], possible_names: list[str]) -> int | None:
    for name in possible_names:
        key = normalise_header(name)

        if key in header_map:
            return header_map[key]

    return None


def cell_value(row: tuple[Any, ...], column_number: int | None) -> Any:
    if column_number is None:
        return None

    index = column_number - 1

    if index >= len(row):
        return None

    return row[index]


def safe_filename(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    value = value.strip("_")
    return value or "new_starter"


def extract_new_starters(
    workbook_path: Path,
    month_start: date,
    month_end: date,
    email_column_name: str | None = None,
    ignore_status: bool = False,
) -> list[NewStarter]:
    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
        keep_vba=True,
    )

    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Could not find sheet '{SHEET_NAME}'. "
            f"Available sheets: {wb.sheetnames}"
        )

    ws = wb[SHEET_NAME]
    header_map = find_header_map(ws)

    candidate_col = require_column(header_map, ["Candidate Name"])
    status_col = require_column(header_map, ["Status"])
    actual_start_col = require_column(header_map, ["Actual Start Date"])

    target_start_col = optional_column(header_map, ["Target Start Date"])
    role_col = optional_column(header_map, ["Role Title (Rate Card)", "Role Title"])
    business_area_col = optional_column(header_map, ["Aviva Business Area"])
    sub_business_area_col = optional_column(header_map, ["Aviva Sub Business Area"])
    capco_lead_col = optional_column(header_map, ["Capco Lead"])
    hiring_manager_col = optional_column(header_map, ["Hiring Manager"])

    recipient_col = None
    if email_column_name:
        recipient_col = optional_column(header_map, [email_column_name])
        if recipient_col is None:
            raise KeyError(
                f"Email column '{email_column_name}' was not found in the sheet."
            )

    starters: list[NewStarter] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        candidate_name = clean_text(cell_value(row, candidate_col))
        status = clean_text(cell_value(row, status_col))
        actual_start_date = parse_excel_date(cell_value(row, actual_start_col))

        if not candidate_name:
            continue

        if actual_start_date is None:
            continue

        if not (month_start <= actual_start_date < month_end):
            continue

        if not ignore_status and "started at aviva" not in normalise_status(status):
            continue

        starter = NewStarter(
            candidate_name=candidate_name,
            role_title=clean_text(cell_value(row, role_col)),
            aviva_business_area=clean_text(cell_value(row, business_area_col)),
            aviva_sub_business_area=clean_text(cell_value(row, sub_business_area_col)),
            capco_lead=clean_text(cell_value(row, capco_lead_col)),
            hiring_manager=clean_text(cell_value(row, hiring_manager_col)),
            status=status,
            actual_start_date=actual_start_date,
            target_start_date=parse_excel_date(cell_value(row, target_start_col)),
            recipient_email=clean_text(cell_value(row, recipient_col)),
        )

        starters.append(starter)

    starters.sort(key=lambda item: (item.actual_start_date, item.candidate_name.casefold()))
    return starters


def first_name(full_name: str) -> str:
    cleaned = clean_text(full_name) or full_name
    return cleaned.split()[0]


def compose_subject(starter: NewStarter) -> str:
    return f"Welcome to the Aviva account, {first_name(starter.candidate_name)}"


def compose_body(starter: NewStarter) -> str:
    role_line = ""
    if starter.role_title:
        role_line = f" It is great to have you joining as {starter.role_title}."

    area_line = ""
    if starter.aviva_business_area or starter.aviva_sub_business_area:
        area_parts = [
            part
            for part in [
                starter.aviva_business_area,
                starter.aviva_sub_business_area,
            ]
            if part
        ]
        area_line = f"\n\nYou are joining the {' / '.join(area_parts)} area."

    lead_line = ""
    if starter.capco_lead:
        lead_line = f"\n\nYour Capco lead is listed as {starter.capco_lead}."

    body = f"""Hi {first_name(starter.candidate_name)},

Welcome to the Aviva account.{role_line} We are pleased to have you on board.

Your recorded start date is {starter.actual_start_date.strftime("%d %B %Y")}.{area_line}{lead_line}

As part of joining the account, please complete the new starter survey using the link below:

{SURVEY_LINK}

This helps us keep the onboarding experience consistent and identify anything we can improve for future joiners.

Thanks,
"""

    return body


def create_email_message(
    starter: NewStarter,
    review_recipient: str | None,
    sender: str | None,
) -> EmailMessage:
    """
    Create an .eml message.

    If review_recipient is provided, every email is addressed to that mailbox.
    This is useful for testing/review before sending to real new starters.

    If review_recipient is not provided, the script uses starter.recipient_email.
    If no recipient email is available, the To field is left blank.
    """
    subject = compose_subject(starter)
    body = compose_body(starter)

    msg = EmailMessage()

    if sender:
        msg["From"] = sender

    msg["To"] = review_recipient or starter.recipient_email or ""
    msg["Subject"] = subject
    msg.set_content(body)

    return msg


def write_audit_csv(output_path: Path, starters: list[NewStarter]) -> None:
    with output_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "candidate_name",
                "recipient_email",
                "role_title",
                "aviva_business_area",
                "aviva_sub_business_area",
                "capco_lead",
                "hiring_manager",
                "status",
                "actual_start_date",
                "target_start_date",
            ],
        )

        writer.writeheader()

        for starter in starters:
            writer.writerow(
                {
                    "candidate_name": starter.candidate_name,
                    "recipient_email": starter.recipient_email or "",
                    "role_title": starter.role_title or "",
                    "aviva_business_area": starter.aviva_business_area or "",
                    "aviva_sub_business_area": starter.aviva_sub_business_area or "",
                    "capco_lead": starter.capco_lead or "",
                    "hiring_manager": starter.hiring_manager or "",
                    "status": starter.status or "",
                    "actual_start_date": starter.actual_start_date.isoformat(),
                    "target_start_date": (
                        starter.target_start_date.isoformat()
                        if starter.target_start_date
                        else ""
                    ),
                }
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Compose deterministic Aviva new starter welcome emails."
    )

    parser.add_argument(
        "--workbook",
        required=True,
        help="Path to aviva_resource_tracker.xlsm.",
    )

    parser.add_argument(
        "--month",
        default=None,
        help="Month to process in YYYY-MM format. Defaults to current month.",
    )

    parser.add_argument(
        "--output-dir",
        default="new_starter_emails",
        help="Directory where .eml files and audit CSV will be written.",
    )

    parser.add_argument(
        "--review-recipient",
        default=None,
        help=(
            "Optional review mailbox. If set, all generated emails are addressed "
            "to this mailbox instead of the new starters."
        ),
    )

    parser.add_argument(
        "--sender",
        default=None,
        help="Optional sender address to include in generated .eml files.",
    )

    parser.add_argument(
        "--email-column",
        default=None,
        help=(
            "Optional column name containing the new starter email address. "
            "The current sheet inspected did not appear to contain an email column."
        ),
    )

    parser.add_argument(
        "--ignore-status",
        action="store_true",
        help=(
            "Use Actual Start Date only and do not require Status to contain "
            "'Started at Aviva'."
        ),
    )

    return parser.parse_args()


def main() -> None:
    args = parse_args()

    workbook_path = Path(args.workbook)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    month_start, month_end = parse_month(args.month)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    starters = extract_new_starters(
        workbook_path=workbook_path,
        month_start=month_start,
        month_end=month_end,
        email_column_name=args.email_column,
        ignore_status=args.ignore_status,
    )

    audit_path = output_dir / "new_starters_audit.csv"
    write_audit_csv(audit_path, starters)

    print("=" * 80)
    print("Aviva new starter email composition")
    print("=" * 80)
    print(f"Workbook:      {workbook_path}")
    print(f"Sheet:         {SHEET_NAME}")
    print(f"Period start:  {month_start.isoformat()}")
    print(f"Period end:    {month_end.isoformat()} exclusive")
    print(f"Starters found: {len(starters)}")
    print(f"Audit CSV:     {audit_path}")

    if not starters:
        print("\nNo new starters found for this period.")
        return

    print("\nGenerated emails:")

    for starter in starters:
        msg = create_email_message(
            starter=starter,
            review_recipient=args.review_recipient,
            sender=args.sender,
        )

        filename = (
            f"{starter.actual_start_date.isoformat()}_"
            f"{safe_filename(starter.candidate_name)}.eml"
        )

        output_path = output_dir / filename
        output_path.write_bytes(bytes(msg))

        print(f"- {starter.candidate_name} -> {output_path}")

    print("\nDone. Open the .eml files in Outlook to review/send.")


if __name__ == "__main__":
    main()